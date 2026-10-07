import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def generate_classroom_excel(classroom, activities, members=None) -> bytes:
    """Generates a master academic plan Excel workbook with formatted styling."""
    wb = Workbook()

    # Style definitions
    font_title = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    font_subtitle = Font(name="Calibri", size=10, italic=True, color="475569")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=10)
    font_bold = Font(name="Calibri", size=10, bold=True)

    fill_header = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    fill_sub_header = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
    fill_completed = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    fill_in_progress = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    fill_planned = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    # Sheet 1: Master Academic Plan
    ws1 = wb.active
    ws1.title = "Master Academic Plan"
    ws1.views.sheetView[0].showGridLines = True

    # Title Banner
    ws1.merge_cells("A1:K1")
    ws1["A1"] = f"ACADEMICFLOW — MASTER PLAN: {classroom.name.upper()}"
    ws1["A1"].font = font_title
    ws1["A1"].alignment = Alignment(horizontal="left", vertical="center")

    ws1.merge_cells("A2:K2")
    ws1["A2"] = (
        f"Join Code: {classroom.join_code} | Department: {classroom.department} | "
        f"Academic Year: {classroom.academic_year} | Exported: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    )
    ws1["A2"].font = font_subtitle
    ws1["A2"].alignment = Alignment(horizontal="left", vertical="center")

    headers = [
        "Activity Code",
        "Faculty",
        "Course",
        "Unit",
        "Activity Topic",
        "Section",
        "Type",
        "Planned Start",
        "Planned End",
        "Status",
        "Completion %",
    ]

    ws1.row_dimensions[4].height = 24
    for col_idx, header in enumerate(headers, 1):
        cell = ws1.cell(row=4, column=col_idx, value=header)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    current_row = 5
    for act in activities:
        ws1.row_dimensions[current_row].height = 19
        row_values = [
            act.activity_id,
            act.faculty or "Unassigned",
            act.course,
            act.unit,
            act.activity_name,
            act.class_section,
            act.activity_type,
            str(act.planned_start),
            str(act.planned_end),
            act.status,
            f"{act.completion_percentage}%",
        ]
        for col_idx, val in enumerate(row_values, 1):
            cell = ws1.cell(row=current_row, column=col_idx, value=val)
            cell.font = font_data
            cell.border = thin_border
            if col_idx in [1, 6, 7, 8, 9]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 11:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.font = font_bold
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

            # Status cell coloring
            if col_idx == 10:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if act.status == "COMPLETED":
                    cell.fill = fill_completed
                    cell.font = Font(name="Calibri", size=10, bold=True, color="166534")
                elif act.status == "IN_PROGRESS":
                    cell.fill = fill_in_progress
                    cell.font = Font(name="Calibri", size=10, bold=True, color="92400E")
                else:
                    cell.fill = fill_planned
            elif current_row % 2 == 0:
                cell.fill = fill_zebra

        current_row += 1

    # Auto-fit columns
    for col in ws1.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Sheet 2: Faculty Roster & Assignments
    if members:
        ws2 = wb.create_sheet(title="Enrolled Faculty")
        ws2.views.sheetView[0].showGridLines = True

        ws2.merge_cells("A1:E1")
        ws2["A1"] = f"FACULTY ROSTER & ENROLLMENT — {classroom.name.upper()}"
        ws2["A1"].font = font_title

        member_headers = ["Teacher Name", "Email", "Assigned Subject", "Assigned Section", "Enrolled Date"]
        ws2.row_dimensions[3].height = 24
        for col_idx, mh in enumerate(member_headers, 1):
            c = ws2.cell(row=3, column=col_idx, value=mh)
            c.font = font_header
            c.fill = fill_sub_header
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = thin_border

        r = 4
        for m in members:
            ws2.row_dimensions[r].height = 19
            t_name = getattr(m.teacher, "name", "Faculty") if hasattr(m, "teacher") and m.teacher else "Faculty"
            t_email = getattr(m.teacher, "email", "—") if hasattr(m, "teacher") and m.teacher else "—"
            m_vals = [
                t_name,
                t_email,
                m.assigned_subject or "All Subjects",
                m.assigned_section or "All Sections",
                m.joined_at.strftime("%Y-%m-%d") if hasattr(m, "joined_at") and m.joined_at else "—",
            ]
            for col_idx, mv in enumerate(m_vals, 1):
                c = ws2.cell(row=r, column=col_idx, value=mv)
                c.font = font_data
                c.border = thin_border
                c.alignment = Alignment(horizontal="center" if col_idx in [4, 5] else "left", vertical="center")
                if r % 2 == 0:
                    c.fill = fill_zebra
            r += 1

        for col in ws2.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws2.column_dimensions[col_letter].width = max(max_len + 4, 15)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
