import logging
from datetime import date
from pathlib import Path
from openpyxl import Workbook, load_workbook

from app.config import ROOT, settings

logger = logging.getLogger("academicflow.excel")
EXCEL_PATH = ROOT / "data" / "master_academic_plan.xlsx"


def ensure_master_excel(activities):
    """Generates master academic plan Excel if absent."""
    if EXCEL_PATH.exists():
        return EXCEL_PATH
    EXCEL_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Academic Plan"
    headers = [
        "Activity ID",
        "Department",
        "Course",
        "Unit",
        "Activity Name",
        "Class Section",
        "Planned Start",
        "Planned End",
        "Actual End",
        "Status",
        "Completion %",
    ]
    ws.append(headers)
    for a in activities:
        ws.append(
            [
                a.activity_id,
                a.department,
                a.course,
                a.unit,
                a.activity_name,
                a.class_section,
                str(a.planned_start),
                str(a.planned_end),
                str(a.actual_end or ""),
                a.status,
                a.completion_percentage,
            ]
        )
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 20
    wb.save(EXCEL_PATH)
    logger.info("Initialized master academic plan Excel at %s", EXCEL_PATH)
    return EXCEL_PATH


def sync_activity_to_excel(activity):
    """Updates the status and completion in the master Excel file."""
    if not EXCEL_PATH.exists():
        return
    try:
        wb = load_workbook(EXCEL_PATH)
        ws = wb.active
        # Find row by Activity ID (column 1)
        updated = False
        for row in range(2, ws.max_row + 1):
            cell_val = str(ws.cell(row=row, column=1).value or "").strip()
            if cell_val == str(activity.activity_id).strip():
                # Column 9: Actual End, Column 10: Status, Column 11: Completion %
                if activity.actual_end:
                    ws.cell(row=row, column=9, value=str(activity.actual_end))
                ws.cell(row=row, column=10, value=activity.status)
                ws.cell(row=row, column=11, value=activity.completion_percentage)
                updated = True
                break
        if updated:
            wb.save(EXCEL_PATH)
            logger.info("Successfully updated Excel row for %s to %s", activity.activity_id, activity.status)
    except Exception as exc:
        logger.error("Failed to sync activity to Excel: %s", exc)
