import hashlib
import io
import json
import re
import uuid
import zipfile
from datetime import date, datetime, time
from pathlib import Path

from fastapi import HTTPException
from openpyxl import load_workbook
from pydantic import ValidationError

from app.config import settings
from app.models import Event, Report
from app.schemas import Extracted
from app.services.audit import record

ALIASES = {
    "activity_description": ["activity", "activity name", "task", "topic", "description"],
    "class_section": ["class", "section", "class section"],
    "event_date": ["date", "execution date", "session date"],
    "faculty": ["faculty", "faculty name", "staff"],
    "course": ["course", "subject", "course name"],
    "department": ["department", "dept"],
    "unit": ["unit", "module"],
    "status": ["status", "completion status"],
    "start_time": ["start time"],
    "end_time": ["end time"],
    "location": ["location", "room"],
    "completion_percentage": ["completion percentage", "completion %"],
}
LOOKUP = {alias: key for key, aliases in ALIASES.items() for alias in aliases}


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for pattern in ["%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d %b %Y", "%d %B %Y"]:
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            pass
    raise ValueError("Date must include a year; supported numeric convention is day/month/year")


def parse_workbook(content):
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024 or len(archive.infolist()) > 1000:
                raise ValueError("Workbook expands beyond permitted size")
            if any("vbaProject" in item.filename for item in archive.infolist()):
                raise ValueError("Macros are not allowed")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise HTTPException(422, "Invalid or unsafe .xlsx workbook") from exc
    results, errors, original, count = [], [], [], 0
    seen = set()
    try:
        for sheet in workbook:
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows, None)
            if not headers:
                continue
            if len(headers) > 100:
                raise HTTPException(422, "At most 100 columns per sheet")
            normalized = [re.sub(r"\s+", " ", str(h or "").strip().lower().replace("_", " ")) for h in headers]
            mapping = [LOOKUP.get(h) for h in normalized]
            recognized = [m for m in mapping if m]
            if len(recognized) != len(set(recognized)):
                raise HTTPException(422, f"Ambiguous duplicate column aliases in sheet {sheet.title}")
            for row_number, row in enumerate(rows, 2):
                if all(v is None or str(v).strip() == "" for v in row):
                    continue
                count += 1
                if count > 5000:
                    raise HTTPException(422, "Maximum 5000 nonblank rows per workbook")
                raw = {
                    "sheet": sheet.title,
                    "row": row_number,
                    "cells": [
                        {"column": i + 1, "header": str(headers[i] or ""), "value": str(v) if v is not None else None}
                        for i, v in enumerate(row)
                    ],
                }
                original.append(raw)
                try:
                    fingerprint = json.dumps([str(v).strip() if v is not None else None for v in row])
                    if fingerprint in seen:
                        raise ValueError("Duplicate row preserved but not processed twice")
                    seen.add(fingerprint)
                    if any(isinstance(v, str) and v.startswith("=") for v in row):
                        raise ValueError("Formula cells are not supported; supply values")
                    fields = {
                        mapping[i]: v.strip() if isinstance(v, str) else v
                        for i, v in enumerate(row)
                        if mapping[i] and v is not None and str(v).strip()
                    }
                    if not fields.get("activity_description"):
                        raise ValueError("An Activity/Topic/Description column and value are required")
                    if "event_date" in fields:
                        fields["event_date"] = parse_date(fields["event_date"])
                    for key in ["start_time", "end_time"]:
                        if key in fields:
                            fields[key] = (
                                fields[key].time()
                                if isinstance(fields[key], datetime)
                                else time.fromisoformat(str(fields[key]))
                            )
                    if "status" in fields:
                        status = str(fields["status"]).strip().upper().replace(" ", "_")
                        fields["status"] = {"DONE": "COMPLETED", "FINISHED": "COMPLETED", "ONGOING": "IN_PROGRESS"}.get(
                            status, status
                        )
                    if "completion_percentage" in fields:
                        fields["completion_percentage"] = float(str(fields["completion_percentage"]).rstrip("%"))
                    fields["source_excerpt"] = json.dumps(raw, ensure_ascii=False)
                    results.append((count, Extracted.model_validate(fields)))
                except (ValueError, ValidationError) as exc:
                    errors.append({"sheet": sheet.title, "row": row_number, "error": str(exc)})
    finally:
        workbook.close()
    if not count:
        raise HTTPException(422, "Workbook contains no report rows")
    return results, errors, original, count


def ingest(db, user, filename, content):
    if not filename or not filename.lower().endswith(".xlsx"):
        raise HTTPException(415, "Only .xlsx files are accepted")
    if len(content) > settings.MAX_UPLOAD_BYTES:
        raise HTTPException(413, "Workbook exceeds upload limit")
    results, errors, original, total = parse_workbook(content)
    safe_name = re.sub(r"[^A-Za-z0-9_.-]", "_", Path(filename.replace("\\", "/")).name)[:120]
    stored_name = f"{uuid.uuid4().hex}.xlsx"
    settings.UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    path = settings.UPLOAD_DIR / stored_name
    path.write_bytes(content)
    report = Report(
        source_type="SPREADSHEET",
        source_file=stored_name,
        raw_content=json.dumps(original, ensure_ascii=False),
        submitted_by=user.id,
        file_metadata={
            "original_name": safe_name,
            "size": len(content),
            "sha256": hashlib.sha256(content).hexdigest(),
            "total_rows": total,
            "processed_rows": len(results),
            "failed_rows": len(errors),
            "errors": errors,
        },
        status="EXTRACTED",
    )
    try:
        db.add(report)
        db.flush()
        record(db, user, "REPORT_RECEIVED", report=report, details=report.file_metadata)
        for row, item in results:
            event = Event(report_id=report.id, source_row=row, **item.model_dump())
            db.add(event)
            db.flush()
            record(
                db,
                user,
                "EXTRACTION_COMPLETED",
                event=event,
                new=item.model_dump(),
                details={"provider": "spreadsheet"},
            )
    except Exception:
        path.unlink(missing_ok=True)
        raise
    return report
